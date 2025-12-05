#!/usr/bin/env python3
"""
Analyze DOCX template structure to understand sections and formatting
"""

from docx import Document
import sys

def analyze_docx(file_path):
    """Analyze DOCX file structure and print detailed information"""
    try:
        doc = Document(file_path)

        print(f"=" * 80)
        print(f"DOCX ANALYSIS: {file_path}")
        print(f"=" * 80)
        print()

        print(f"Total Paragraphs: {len(doc.paragraphs)}")
        print(f"Total Tables: {len(doc.tables)}")
        print()

        print("-" * 80)
        print("DOCUMENT STRUCTURE:")
        print("-" * 80)

        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:  # Only show non-empty paragraphs
                style = para.style.name
                is_bold = para.runs[0].bold if para.runs else False
                font_size = para.runs[0].font.size if para.runs and para.runs[0].font.size else "Default"

                # Detect section headers (usually bold, larger font, or specific style)
                is_likely_header = (
                    is_bold or
                    'Heading' in style or
                    text.isupper() or
                    len(text) < 50  # Short text often indicates headers
                )

                marker = "📌 HEADER" if is_likely_header else "   "
                print(f"{i:3d} | {marker} | {style:20s} | {text[:60]}")

        # Analyze tables if present
        if doc.tables:
            print()
            print("-" * 80)
            print("TABLES FOUND:")
            print("-" * 80)
            for i, table in enumerate(doc.tables):
                print(f"\nTable {i+1}: {len(table.rows)} rows x {len(table.columns)} columns")
                # Show first few cells
                for row_idx, row in enumerate(table.rows[:3]):
                    cells_text = [cell.text[:30] for cell in row.cells]
                    print(f"  Row {row_idx}: {cells_text}")

        return doc

    except Exception as e:
        print(f"Error analyzing DOCX: {e}")
        import traceback
        traceback.print_exc()
        return None

if __name__ == "__main__":
    file_path = "Henry Jung_CV_Imperial Brands PLC.docx"
    analyze_docx(file_path)
